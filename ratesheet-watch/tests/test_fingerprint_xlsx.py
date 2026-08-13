from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill
import fingerprint

FIXTURE = Path("/Users/trungthach/IdeaProjects/packs/loan/src/test/resources/ratesheets/CW_2026_06_09.xlsx")
GOLDEN = Path(__file__).parent / "golden" / "cw_20260609.json"

def test_xlsx_extract_shape_and_golden():
    fp = fingerprint.extract(FIXTURE)
    assert fp["format"] == "xlsx"
    assert fp["pages"] >= 1
    assert fp["banners"] and fp["structure"]
    if not GOLDEN.exists():
        GOLDEN.write_text(fingerprint.to_json(fp))
    assert fingerprint.to_json(fp) == GOLDEN.read_text()

def test_banner_cell_by_fill(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "6.125"                                     # numeric, plain
    ws["A2"] = ".375 Price Improvement for loans >$350K"   # filled banner
    ws["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B5"] = "30 Year Fixed"                             # plain short-ish text
    p = tmp_path / "t.xlsx"
    wb.save(p)
    fp = fingerprint.extract(p)
    assert any("Price Improvement" in b for b in fp["banners"])
    assert any("#" in s for s in fp["structure"])          # 6.125 masked

def test_embedded_images_hashed(tmp_path):
    import zipfile, shutil
    src = tmp_path / "img.xlsx"
    wb = openpyxl.Workbook(); wb.save(src)
    # inject a fake media file into the zip (openpyxl image API needs Pillow)
    with zipfile.ZipFile(src, "a") as z:
        z.writestr("xl/media/image1.png", b"\x89PNGfakebytes")
    fp = fingerprint.extract(src)
    assert len(fp["images"]) == 1 and len(fp["images"][0]) == 40  # sha1 hex
    imgs = fingerprint.embedded_images(src)
    assert imgs[0][0] == "xl/media/image1.png" and imgs[0][1].startswith(b"\x89PNG")

def test_xls_unextractable(tmp_path):
    """Old binary .xls files cannot be read by openpyxl; mark unextractable."""
    p = tmp_path / "old.xls"
    # Write the OLE2 header that identifies .xls files
    p.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 100)
    fp = fingerprint.extract(p)
    assert fp["format"] == "xls"
    assert fp["unextractable"] is True
    assert fp["banners"] == []
    assert fp["structure"] == []

def test_numeric_merged_cell_is_structure_not_banner(tmp_path):
    """Numeric content in wide merged cells (≥4 cols) goes to structure, not banners."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:F1")
    ws["A1"] = "6.125 6.250 6.375"  # numeric line spanning 6 columns
    p = tmp_path / "merged_numeric.xlsx"
    wb.save(p)
    fp = fingerprint.extract(p)
    # Should not appear in banners (numeric content)
    assert not any("6.125" in b for b in fp["banners"])
    # Should appear in structure with masked numbers
    assert any("#" in s for s in fp["structure"])
